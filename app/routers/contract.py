from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from pydantic import BaseModel
from typing import Optional
from sqlalchemy.ext.asyncio import AsyncSession

from app.utils.dependencies import get_current_user
from app.models.user import User
from app.openai_client import generate_contract_with_model, review_contract_with_model

from app.database import get_db
from app.models.contract import Contract
from app.models.review import Review
from app.models.activity import Activity

from sqlalchemy import select
from app.models.contract import Contract

router = APIRouter(tags=["Contract"])


# -----------------------------
# Request Models
# -----------------------------
class ContractGenerateRequest(BaseModel):
    contract_type: str
    details: dict
    language: str = "English"
    tone: str = "formal"
    context: str = ""

    def is_custom_contract(self):
        return str(self.contract_type).lower().strip() == "custom"

    def get_custom_description(self):
        """
        For Custom Contract simple mode:
        Frontend should send:
        {
            "contract_type": "custom",
            "details": {
                "custom_description": "User's full contract description..."
            }
        }

        To be safe, we also accept a few alternative keys.
        """
        if not isinstance(self.details, dict):
            return ""

        return (
            self.details.get("custom_description")
            or self.details.get("description")
            or self.details.get("details")
            or self.details.get("contract_description")
            or ""
        )

    def get_answers(self):
        """
        Keeps the previous behavior for all normal contract types.
        For custom contract, it normalizes the user's description so the AI layer
        always receives a clear custom_description field.
        """
        if not isinstance(self.details, dict):
            return {}

        if self.is_custom_contract():
            custom_description = self.get_custom_description()

            normalized_details = dict(self.details)
            normalized_details["custom_description"] = custom_description

            return normalized_details

        return self.details


class ContractReviewRequest(BaseModel):
    contract_text: str
    language: str = "English"


# -----------------------------
# Generate Contract
# -----------------------------
@router.post("/generate")
async def generate(
    data: ContractGenerateRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    # -----------------------------
    # Custom Contract Simple Mode
    # -----------------------------
    # This does NOT create a new endpoint.
    # It only allows the existing /contracts/generate endpoint
    # to generate a custom contract from user's free-text description.
    if data.is_custom_contract():
        custom_description = data.get_custom_description()

        if not custom_description or not str(custom_description).strip():
            raise HTTPException(
                status_code=400,
                detail="custom_description is required for custom contract."
            )

        # If context is empty, use the custom description as context too.
        # This helps the AI generation layer understand the user's intent
        # without breaking the previous function signature.
        if not data.context or not str(data.context).strip():
            data.context = str(custom_description).strip()

    # ----- منطق اصلی تولید قرارداد (بدون تغییر) -----
    result = await generate_contract_with_model(
        contract_type=data.contract_type,
        language=data.language,
        tone=data.tone,
        context=data.context,
        answers=data.get_answers(),
    )

    # -----------------------------
    # Save Contract + Activity
    # -----------------------------
    contract_title = f"{data.contract_type} Contract"

    if data.is_custom_contract():
        contract_title = "Custom Contract"

    new_contract = Contract(
        user_id=current_user.id,
        title=contract_title,
        contract_type=data.contract_type,
        language=data.language,
        content=result,
    )

    db.add(new_contract)
    await db.commit()
    await db.refresh(new_contract)

    activity_message = f"Generated {data.contract_type} contract"

    if data.is_custom_contract():
        activity_message = "Generated custom contract"

    activity = Activity(
        user_id=current_user.id,
        event_type="contract_generate",
        message=activity_message
    )
    db.add(activity)
    await db.commit()

    return {
        "contract": result,
        "contract_id": new_contract.id
    }


# -----------------------------
# Review Contract (Text)
# -----------------------------
@router.post("/review")
async def review(
    data: ContractReviewRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    # ----- منطق اصلی بررسی قرارداد (بدون تغییر) -----
    result = await review_contract_with_model(
        text=data.contract_text,
        language=data.language,
    )

    # -----------------------------
    # Save Review + Activity
    # -----------------------------
    new_review = Review(
        user_id=current_user.id,
        original_text=data.contract_text,
        review_output=str(result),
        language=data.language,
    )

    db.add(new_review)
    await db.commit()
    await db.refresh(new_review)

    activity = Activity(
        user_id=current_user.id,
        event_type="contract_review",
        message="Reviewed a contract text"
    )
    db.add(activity)
    await db.commit()

    return result


# -----------------------------
# Review Contract (File Upload)
# -----------------------------
@router.post("/review-file")
async def review_file(
    file: UploadFile = File(...),
    language: str = Form(default="English"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    filename = file.filename.lower()

    # ----------- Extract Text (بدون تغییر منطق) -----------
    if filename.endswith(".txt"):
        content = await file.read()
        text = content.decode("utf-8", errors="ignore")

    elif filename.endswith(".pdf"):
        try:
            import pdfplumber, io
            content = await file.read()
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        except ImportError:
            raise HTTPException(status_code=500, detail="pdfplumber not installed")

    elif filename.endswith(".docx"):
        try:
            import docx, io
            content = await file.read()
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            raise HTTPException(status_code=500, detail="python-docx not installed")

    else:
        raise HTTPException(status_code=400, detail="Unsupported file type. Use PDF, DOCX, or TXT.")

    if not text.strip():
        raise HTTPException(status_code=400, detail="Could not extract text from file.")

    # ----------- منطق اصلی بررسی فایل (بدون تغییر) -----------
    result = await review_contract_with_model(
        text=text,
        language=language
    )

    # -----------------------------
    # Save Review + Activity
    # -----------------------------
    new_review = Review(
        user_id=current_user.id,
        original_text=text,
        review_output=str(result),
        language=language,
    )

    db.add(new_review)
    await db.commit()
    await db.refresh(new_review)

    activity = Activity(
        user_id=current_user.id,
        event_type="contract_review_file",
        message=f"Reviewed contract file: {file.filename}"
    )
    db.add(activity)
    await db.commit()

    return result


@router.post("/upload")
async def upload_contract(
    file: UploadFile = File(...),
    title: str = Form(default="Uploaded Contract"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    filename = file.filename.lower()
    content = await file.read()

    # Extract text from different file types
    if filename.endswith(".txt"):
        text = content.decode("utf-8", errors="ignore")

    elif filename.endswith(".pdf"):
        import pdfplumber, io
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)

    elif filename.endswith(".docx"):
        import docx, io
        doc = docx.Document(io.BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs)

    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Use PDF, DOCX, or TXT.",
        )

    # Save contract to database
    new_contract = Contract(
        user_id=current_user.id,
        title=title,
        contract_type=file.content_type,
        content=text,
        language="English",
    )

    db.add(new_contract)
    await db.commit()
    await db.refresh(new_contract)

    return {
        "id": new_contract.id,
        "title": new_contract.title,
        "type": new_contract.contract_type,
        "status": "uploaded",
    }


async def get_contract_text(contract_id: int, db: AsyncSession):
    stmt = select(Contract).where(Contract.id == contract_id)
    result = await db.execute(stmt)
    contract = result.scalar_one_or_none()

    if not contract:
        raise HTTPException(status_code=404, detail="Contract not found")

    return contract.content
