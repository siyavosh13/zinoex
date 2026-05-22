import io
import re
import unicodedata


def clean_text(text: str) -> str:
    if not text:
        return ""

    # normalize unicode
    text = unicodedata.normalize("NFKC", text)

    # remove control characters
    text = re.sub(r"[\x00-\x1F\x7F]", " ", text)

    # remove extra spaces
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def extract_pdf(contents: bytes) -> str:
    text_parts = []

    # ---------- FIRST TRY (pdfplumber) ----------
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(contents)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()

                if page_text and page_text.strip():
                    text_parts.append(page_text)

    except Exception as e:
        print("pdfplumber failed:", e)

    # ---------- FALLBACK (pypdf) ----------
    if not text_parts:
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(contents))

            for page in reader.pages:
                page_text = page.extract_text()

                if page_text and page_text.strip():
                    text_parts.append(page_text)

        except Exception as e:
            print("pypdf fallback failed:", e)

    text = "\n".join(text_parts)

    return clean_text(text)


def extract_docx(contents: bytes) -> str:
    from docx import Document

    text_parts = []

    try:
        doc = Document(io.BytesIO(contents))

        # paragraphs
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                text_parts.append(p.text)

        # tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        text_parts.append(cell.text)

    except Exception as e:
        print("docx extraction failed:", e)

    text = "\n".join(text_parts)

    return clean_text(text)


def extract_txt(contents: bytes) -> str:
    try:
        text = contents.decode("utf-8", errors="ignore")
    except Exception:
        text = contents.decode("latin-1", errors="ignore")

    return clean_text(text)


def extract_file(filename: str, contents: bytes) -> str:
    filename = filename.lower()

    if filename.endswith(".pdf"):
        return extract_pdf(contents)

    elif filename.endswith(".docx"):
        return extract_docx(contents)

    elif filename.endswith(".txt"):
        return extract_txt(contents)

    else:
        raise ValueError("Unsupported file type")
